import pytest
from docx import Document as WordDocument

from app.ingestion.document_loader import DocumentLoader


@pytest.fixture
def loader():
    return DocumentLoader()


def test_load_txt_preserves_source_metadata(loader, tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Enterprise RAG project", encoding="utf-8")

    documents = loader.load_file(path)

    assert len(documents) == 1
    assert documents[0].page_content == "Enterprise RAG project"
    assert documents[0].metadata["source"] == "notes.txt"
    assert documents[0].metadata["file_type"] == "txt"


def test_load_markdown(loader, tmp_path):
    path = tmp_path / "guide.md"
    path.write_text("# Retrieval\nHybrid search combines dense and sparse retrieval.")

    documents = loader.load_file(path)

    assert len(documents) == 1
    assert "Hybrid search" in documents[0].page_content
    assert documents[0].metadata["file_type"] == "md"


def test_load_csv_creates_one_document_per_nonempty_row(loader, tmp_path):
    path = tmp_path / "customers.csv"
    path.write_text(
        "name,city\nRahul,Hyderabad\nAnita,Adoni\n,,\n",
        encoding="utf-8",
    )

    documents = loader.load_file(path)

    assert len(documents) == 2
    assert "name: Rahul" in documents[0].page_content
    assert "city: Hyderabad" in documents[0].page_content
    assert documents[0].metadata["row"] == 2
    assert documents[1].metadata["row"] == 3
    assert documents[0].metadata["file_type"] == "csv"


def test_load_html_ignores_script_and_style(loader, tmp_path):
    path = tmp_path / "page.html"
    path.write_text(
        """
        <html>
          <body>
            <h1>Company Policy</h1>
            <p>Employees must protect confidential data.</p>
            <script>fake_secret = 'ignore me'</script>
            <style>.hidden { display: none; }</style>
          </body>
        </html>
        """,
        encoding="utf-8",
    )

    documents = loader.load_file(path)
    content = documents[0].page_content

    assert "Company Policy" in content
    assert "protect confidential data" in content
    assert "ignore me" not in content
    assert "display: none" not in content


def test_load_docx_includes_paragraphs_and_tables(loader, tmp_path):
    path = tmp_path / "report.docx"

    word_document = WordDocument()
    word_document.add_paragraph("Quarterly Enterprise Report")

    table = word_document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Revenue"
    table.cell(0, 1).text = "500000"
    word_document.save(path)

    documents = loader.load_file(path)
    content = documents[0].page_content

    assert "Quarterly Enterprise Report" in content
    assert "Revenue" in content
    assert "500000" in content
    assert documents[0].metadata["file_type"] == "docx"


def test_unsupported_extension_raises_value_error(loader, tmp_path):
    path = tmp_path / "notes.xyz"
    path.write_text("Some content", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file type"):
        loader.load_file(path)


def test_missing_file_raises_file_not_found(loader, tmp_path):
    with pytest.raises(FileNotFoundError):
        loader.load_file(tmp_path / "missing.txt")


def test_empty_text_file_raises_value_error(loader, tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("", encoding="utf-8")

    with pytest.raises(ValueError, match="No extractable text"):
        loader.load_file(path)