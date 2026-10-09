import csv
from html.parser import HTMLParser
from pathlib import Path

from docx import Document as WordDocument
from langchain_core.documents import Document
from pypdf import PdfReader


class _HTMLTextExtractor(HTMLParser):
    """Extract readable text while ignoring scripts and styles."""

    IGNORED_TAGS = {"script", "style", "noscript"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in self.IGNORED_TAGS:
            self._ignored_depth += 1
        elif tag.lower() in {"p", "div", "br", "li", "tr", "h1", "h2", "h3"}:
            if self._ignored_depth == 0:
                self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in self.IGNORED_TAGS and self._ignored_depth:
            self._ignored_depth -= 1
        elif tag.lower() in {"p", "div", "li", "tr", "h1", "h2", "h3"}:
            if self._ignored_depth == 0:
                self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._ignored_depth == 0 and data.strip():
            self.parts.append(data.strip())

    def get_text(self) -> str:
        return "\n".join(part for part in self.parts if part.strip())


class DocumentLoader:
    """Load supported files into standardized LangChain documents."""

    SUPPORTED_EXTENSIONS = {
        ".pdf",
        ".txt",
        ".docx",
        ".md",
        ".markdown",
        ".csv",
        ".html",
        ".htm",
    }

    def load_file(self, file_path: str | Path) -> list[Document]:
        path = Path(file_path)

        if not path.is_file():
            raise FileNotFoundError(f"Document not found: {path}")

        extension = path.suffix.lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            supported = ", ".join(sorted(self.SUPPORTED_EXTENSIONS))
            raise ValueError(
                f"Unsupported file type: {extension or '(no extension)'}. "
                f"Supported extensions: {supported}"
            )

        loaders = {
            ".pdf": self._load_pdf,
            ".txt": self._load_text,
            ".md": self._load_text,
            ".markdown": self._load_text,
            ".docx": self._load_docx,
            ".csv": self._load_csv,
            ".html": self._load_html,
            ".htm": self._load_html,
        }

        documents = loaders[extension](path)

        documents = [
            document
            for document in documents
            if document.page_content.strip()
        ]

        if not documents:
            raise ValueError(f"No extractable text found in: {path.name}")

        return documents

    @staticmethod
    def _load_pdf(path: Path) -> list[Document]:
        reader = PdfReader(str(path))
        documents = []

        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""

            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": path.name,
                        "page": page_number,
                        "file_type": "pdf",
                    },
                )
            )

        return documents

    @staticmethod
    def _load_text(path: Path) -> list[Document]:
        file_type = path.suffix.lower().lstrip(".")
        text = path.read_text(encoding="utf-8-sig")

        return [
            Document(
                page_content=text,
                metadata={
                    "source": path.name,
                    "file_type": file_type,
                },
            )
        ]

    @staticmethod
    def _load_docx(path: Path) -> list[Document]:
        word_document = WordDocument(str(path))
        parts = []

        for paragraph in word_document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        # Include table content because important information often
        # appears in Word tables rather than paragraphs.
        for table_index, table in enumerate(word_document.tables, start=1):
            for row in table.rows:
                values = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]
                if any(values):
                    parts.append(" | ".join(values))

        return [
            Document(
                page_content="\n".join(parts),
                metadata={
                    "source": path.name,
                    "file_type": "docx",
                },
            )
        ]

    @staticmethod
    def _load_csv(path: Path) -> list[Document]:
        documents = []

        with path.open(
            "r",
            encoding="utf-8-sig",
            newline="",
        ) as file:
            reader = csv.DictReader(file)

            if not reader.fieldnames:
                return []

            columns = [
                column.strip() if column else ""
                for column in reader.fieldnames
            ]

            for row_number, row in enumerate(reader, start=2):
                lines = []

                for original_column, cleaned_column in zip(
                    reader.fieldnames, columns
                ):
                    value = row.get(original_column)

                    if value is None or not str(value).strip():
                        continue

                    lines.append(f"{cleaned_column}: {str(value).strip()}")

                if not lines:
                    continue

                documents.append(
                    Document(
                        page_content="\n".join(lines),
                        metadata={
                            "source": path.name,
                            "file_type": "csv",
                            "row": row_number,
                            "columns": ", ".join(columns),
                        },
                    )
                )

        return documents

    @staticmethod
    def _load_html(path: Path) -> list[Document]:
        html = path.read_text(encoding="utf-8-sig")
        parser = _HTMLTextExtractor()
        parser.feed(html)
        parser.close()

        return [
            Document(
                page_content=parser.get_text(),
                metadata={
                    "source": path.name,
                    "file_type": path.suffix.lower().lstrip("."),
                },
            )
        ]